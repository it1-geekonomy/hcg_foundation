"""Extract plain text from PDF / DOCX / PPTX for chatbot knowledge."""

from __future__ import annotations

import logging
import re
from pathlib import Path

logger = logging.getLogger(__name__)

_WS = re.compile(r"[ \t]+")
_BLANK = re.compile(r"\n{3,}")


def _normalize(text: str) -> str:
    text = text.replace("\x00", " ")
    text = _WS.sub(" ", text)
    text = _BLANK.sub("\n\n", text)
    return text.strip()


def extract_pdf(path: Path) -> str:
    parts: list[str] = []

    # Prefer PyMuPDF when available (better layout / scanned text layers)
    try:
        import fitz  # type: ignore

        doc = fitz.open(str(path))
        try:
            for page in doc:
                parts.append(page.get_text("text") or "")
        finally:
            doc.close()
        text = _normalize("\n".join(parts))
        if text:
            return text
    except Exception as exc:
        logger.debug("pymupdf failed for %s: %s", path.name, exc)

    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        for page in reader.pages:
            parts.append(page.extract_text() or "")
        return _normalize("\n".join(parts))
    except Exception as exc:
        logger.warning("PDF extract failed for %s: %s", path.name, exc)
        return ""


def extract_docx(path: Path) -> str:
    try:
        from docx import Document  # type: ignore

        doc = Document(str(path))
        parts: list[str] = []
        for para in doc.paragraphs:
            t = (para.text or "").strip()
            if t:
                parts.append(t)
        for table in doc.tables:
            for row in table.rows:
                cells = [
                    (c.text or "").strip() for c in row.cells if (c.text or "").strip()
                ]
                if cells:
                    parts.append(" | ".join(cells))
        return _normalize("\n".join(parts))
    except Exception as exc:
        logger.warning("DOCX extract failed for %s: %s", path.name, exc)
        return ""


def extract_pptx(path: Path) -> str:
    try:
        from pptx import Presentation  # type: ignore

        prs = Presentation(str(path))
        parts: list[str] = []
        for i, slide in enumerate(prs.slides, 1):
            slide_bits: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    t = (shape.text or "").strip()
                    if t:
                        slide_bits.append(t)
            if slide_bits:
                parts.append(f"Slide {i}:\n" + "\n".join(slide_bits))
        return _normalize("\n\n".join(parts))
    except Exception as exc:
        logger.warning("PPTX extract failed for %s: %s", path.name, exc)
        return ""


def extract_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".pptx":
        return extract_pptx(path)
    if suffix in {".txt", ".md"}:
        return _normalize(path.read_text(encoding="utf-8", errors="ignore"))
    return ""
